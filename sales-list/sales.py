"""営業リスト (sales_list.xlsx) の作成・更新スクリプト。

使い方:
  python sales.py init                                   # 新規作成（既存があれば何もしない）
  python sales.py add    <タブ> <名前> 確度=A 所在=鳥取市 ...  # 行を追加
  python sales.py update <タブ> <名前> 確度=B 次のアクション=...  # 既存行を更新

update では変更前の値を「[タイムスタンプ] 旧: 値」としてセルのコメントに追記し、
「更新履歴」タブにも1行残す。最終更新日は自動で更新される。
"""
import sys
from datetime import datetime
from zoneinfo import ZoneInfo
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

FILE = Path(__file__).with_name("sales_list.xlsx")
TABS = ["ゴールド", "ゴールド研修", "ボードゲーム交流会", "企業研修", "HP", "ライティング", "鳥取ツアー", "理学療法"]
COLUMNS = ["名前", "確度", "所在", "連絡先", "ステータス", "次のアクション", "次回期日", "最終更新日", "メモ"]
WIDTHS = [20, 7, 16, 22, 14, 30, 12, 18, 40]
RANKS = {"A": ("受注目前（80%〜）", "F8CBAD"), "B": ("前向き検討（50%〜）", "FFE699"),
         "C": ("関心あり（20%〜）", "C6E0B4"), "D": ("見込み薄・接点のみ", "D9D9D9"),
         "済": ("受注済み", "9BC2E6"), "失注": ("失注", "BFBFBF")}
STATUSES = ["未接触", "アプローチ中", "提案中", "見積提出", "交渉中", "受注", "失注", "保留"]
HISTORY = "更新履歴"
FONT = "Arial"
THIN = Side(style="thin", color="BFBFBF")


def now():
    return datetime.now(ZoneInfo("Asia/Tokyo")).strftime("%Y-%m-%d %H:%M")


def style_header(ws, headers, widths):
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(headers)).column_letter}1000"


def build():
    wb = Workbook()
    guide = wb.active
    guide.title = "使い方"
    lines = [
        ("営業リスト 使い方", True),
        ("", False),
        ("■ 更新の流れ", True),
        ("チャットに「〇〇（タブ名）の△△さん、確度Bに。来週見積送る」のように進捗を送る → このファイルが更新されます。", False),
        ("変更前の値はセルのコメントに「[日時] 旧: 値」で残り、『更新履歴』タブにも記録されます。", False),
        ("手で直接編集してもOK（その場合コメントと履歴は残りません）。", False),
        ("", False),
        ("■ 必須項目：名前・確度・所在", True),
        ("", False),
        ("■ 確度の基準", True),
    ] + [(f"{k}：{v[0]}", False) for k, v in RANKS.items()] + [
        ("", False),
        ("■ ステータス候補", True),
        ("／".join(STATUSES), False),
        ("", False),
        ("※『ゴールド』タブ2行目は記入例です。不要になったら削除してください。", False),
    ]
    for r, (text, bold) in enumerate(lines, 1):
        c = guide.cell(row=r, column=1, value=text)
        c.font = Font(name=FONT, bold=bold, size=14 if r == 1 else 11)
        key = text.split("：")[0]
        if key in RANKS and not bold:
            c.fill = PatternFill("solid", fgColor=RANKS[key][1])
    guide.column_dimensions["A"].width = 110

    for name in TABS:
        ws = wb.create_sheet(name)
        style_header(ws, COLUMNS, WIDTHS)
        dv_rank = DataValidation(type="list", formula1='"' + ",".join(RANKS) + '"', allow_blank=True)
        dv_status = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
        ws.add_data_validation(dv_rank)
        ws.add_data_validation(dv_status)
        dv_rank.add("B2:B1000")
        dv_status.add("E2:E1000")
        for k, (_, color) in RANKS.items():
            ws.conditional_formatting.add(
                "B2:B1000", FormulaRule(formula=[f'$B2="{k}"'], fill=PatternFill("solid", fgColor=color)))
        # 期日超過は赤字
        ws.conditional_formatting.add(
            "G2:G1000", FormulaRule(formula=['AND($G2<>"",$G2<TODAY())'], font=Font(color="C00000", bold=True)))

    h = wb.create_sheet(HISTORY)
    style_header(h, ["日時", "タブ", "名前", "項目", "変更前", "変更後"], [18, 16, 20, 16, 30, 30])

    wb.save(FILE)
    add_row("ゴールド", "（例）山田太郎", {"確度": "B", "所在": "鳥取市", "連絡先": "090-xxxx-xxxx",
                                         "ステータス": "提案中", "次のアクション": "見積を送付する",
                                         "次回期日": "2026-10-09", "メモ": "記入例。削除してOK"})


def cell_style(c):
    c.font = Font(name=FONT)
    c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    c.alignment = Alignment(vertical="top", wrap_text=True)


def find_row(ws, name):
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == name:
            return r
    return None


def parse_value(col, v):
    if col == "次回期日" and v:
        try:
            return datetime.strptime(v, "%Y-%m-%d")
        except ValueError:
            pass
    return v


def add_row(tab, name, fields):
    wb = load_workbook(FILE)
    ws = wb[tab]
    if find_row(ws, name):
        sys.exit(f"{tab} に {name} は既にあります。update を使ってください。")
    for req in ("確度", "所在"):
        if not fields.get(req):
            sys.exit(f"必須項目「{req}」がありません。")
    r = ws.max_row + 1
    while r > 2 and ws.cell(row=r - 1, column=1).value is None:
        r -= 1
    fields = {"名前": name, **fields, "最終更新日": now()}
    for i, col in enumerate(COLUMNS, 1):
        c = ws.cell(row=r, column=i, value=parse_value(col, fields.get(col)))
        cell_style(c)
        if col == "次回期日":
            c.number_format = "yyyy-mm-dd"
    log(wb, tab, name, "新規追加", "", ", ".join(f"{k}={v}" for k, v in fields.items() if k != "名前"))
    wb.save(FILE)


def update_row(tab, name, fields):
    wb = load_workbook(FILE)
    ws = wb[tab]
    r = find_row(ws, name)
    if not r:
        sys.exit(f"{tab} に {name} が見つかりません。add を使ってください。")
    ts = now()
    for col, new in fields.items():
        if col not in COLUMNS:
            sys.exit(f"不明な項目: {col}（使える項目: {', '.join(COLUMNS)}）")
        c = ws.cell(row=r, column=COLUMNS.index(col) + 1)
        old = c.value
        old_s = old.strftime("%Y-%m-%d") if isinstance(old, datetime) else ("" if old is None else str(old))
        if old_s == new:
            continue
        note = f"[{ts}] 旧: {old_s or '（空欄）'}"
        prev = c.comment.text + "\n" if c.comment else ""
        c.value = parse_value(col, new)
        c.comment = Comment(prev + note, "営業リスト")
        c.comment.width, c.comment.height = 300, 120
        cell_style(c)
        if col == "次回期日":
            c.number_format = "yyyy-mm-dd"
        log(wb, tab, name, col, old_s, new)
    ws.cell(row=r, column=COLUMNS.index("最終更新日") + 1, value=ts)
    wb.save(FILE)


def log(wb, tab, name, col, old, new):
    h = wb[HISTORY]
    r = h.max_row + 1
    for i, v in enumerate([now(), tab, name, col, old, new], 1):
        cell_style(h.cell(row=r, column=i, value=v))


def parse_fields(args):
    out = {}
    for a in args:
        k, _, v = a.partition("=")
        out[k] = v
    return out


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else "init"
    if cmd == "init":
        if FILE.exists():
            sys.exit(f"{FILE.name} は既にあります。")
        build()
    elif cmd in ("add", "update"):
        tab, name = sys.argv[2], sys.argv[3]
        (add_row if cmd == "add" else update_row)(tab, name, parse_fields(sys.argv[4:]))
    else:
        sys.exit(__doc__)
    print("OK:", FILE)
